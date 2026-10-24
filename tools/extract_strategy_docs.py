#!/usr/bin/env python3
"""Extract text from .docx strategy documents"""
import sys
from pathlib import Path
from docx import Document

def extract_docx(docx_path):
    """Extract all text from a .docx file"""
    doc = Document(docx_path)
    full_text = []

    for para in doc.paragraphs:
        full_text.append(para.text)

    # Also extract from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text.append(cell.text)

    return '\n'.join(full_text)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python extract_strategy_docs.py <docx_file>")
        sys.exit(1)

    docx_file = sys.argv[1]
    text = extract_docx(docx_file)
    print(text)
